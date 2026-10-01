"""Сверка РСВ ООО «Легкие деньки» за 1 квартал и полугодие 2026 г.: первичные и уточнённые расчёты по сотрудникам."""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "reports/svod_2026/Сверка_РСВ_2026.xlsx"

thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DDE7F3")
TOTAL = PatternFill("solid", fgColor="F2F2F2")
BAD = PatternFill("solid", fgColor="F8D7DA")
OK = PatternFill("solid", fgColor="D4EDDA")
NUM = '#,##0.00;[Red]-#,##0.00;"–"'

COLS = ["Сотрудник", "Месяц",
        "Выплаты: {a}", "Выплаты: {b}", "Разница выплат ({b} − {a})",
        "Взносы: {a}", "Взносы: {b}", "Разница взносов ({b} − {a})",
        "Проверка: взносы {b} − 30% от выплат", "Комментарий"]
WIDTHS = [30, 11, 16, 16, 16, 16, 16, 16, 16, 60]

wb = Workbook()


def sheet(title, header_lines, a, b):
    ws = wb.create_sheet(title)
    r = 1
    for i, line in enumerate(header_lines):
        ws.cell(r, 1, line).font = Font(bold=(i == 0), size=13 if i == 0 else 10)
        r += 1
    r += 1
    for c, name in enumerate(COLS, 1):
        cell = ws.cell(r, c, name.format(a=a, b=b))
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r].height = 62
    for c, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(r + 1, 3)
    return ws, r + 1


def row(ws, r, name, month, pay_a, pay_b, vz_a, vz_b, note="", fill=None, bold=False):
    vals = [name, month, pay_a, pay_b, f"=D{r}-C{r}", vz_a, vz_b, f"=G{r}-F{r}",
            f"=G{r}-ROUND(D{r}*0.3,2)", note]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = BORDER
        if 3 <= c <= 9:
            cell.number_format = NUM
        else:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            cell.fill = fill
        if bold:
            cell.font = Font(bold=True)
    for L in "EHI":
        ws.conditional_formatting.add(f"{L}{r}", CellIsRule(operator="notBetween", formula=["-0.005", "0.005"], fill=BAD))
        ws.conditional_formatting.add(f"{L}{r}", CellIsRule(operator="between", formula=["-0.005", "0.005"], fill=OK))


def person(ws, r, name, months, notes=None):
    """months: [(месяц, выплаты A, выплаты B, взносы A, взносы B)]; добавляет строку «Итого»."""
    notes = notes or {}
    start = r
    for m in months:
        row(ws, r, name if r == start else "", *m, note=notes.get(m[0], ""))
        r += 1
    row(ws, r, f"Итого {name.split()[0]}", "",
        *[f"=SUM({L}{start}:{L}{r - 1})" for L in "CDFG"], note=notes.get("итого", ""), fill=TOTAL, bold=True)
    return r + 1, r


def text(ws, r, title, lines):
    ws.cell(r, 1, title).font = Font(bold=True)
    for t in lines:
        r += 1
        ws.cell(r, 1, t)
    return r + 2


# ----------------------------------------------------------------- 1 квартал
ws, r = sheet("1 квартал", [
    "РСВ за 1 квартал 2026 г.: первичный и уточнённый — по сотрудникам",
    "ООО «Легкие деньки», ИНН 3663129244, тариф 01 (30%)",
    "Первичный: корр. 0 от 01.06.2026, КПП 366301001, ИФНС 3663.   Уточнённый: корр. 1 от 08.09.2026, КПП 366201001, ИФНС 3662.",
], "первичный", "уточнённый")

r, tot_l = person(ws, r, "Лубкова Наталья Ивановна", [
    ("Январь", 27093.00, 27093.00, 8127.90, 8127.90),
    ("Февраль", 27093.00, 27093.00, 8127.90, 8127.90),
    ("Март", 20101.26, 27093.00, 6030.38, 8127.90),
], {"Март": "Расхождение: в первичном март 20 101,26 ₽, в уточнённом — полный МРОТ 27 093,00 ₽."})
r, tot_b = person(ws, r, "Баскова Неонила Павловна", [
    ("Январь", 0, 0, 0, 0),
], {"Январь": "В обоих расчётах — код НР, выплат нет."})

row(ws, r, "ИТОГО по разделу 3", "1 кв.", *[f"={L}{tot_l}+{L}{tot_b}" for L in "CDFG"], fill=TOTAL, bold=True)
t3 = r
r += 1
row(ws, r, "Раздел 1 (подр. 1 стр. 030 / разд. 1 стр. 030)", "1 кв.", 74287.26, 81279.00, 22286.18, 24383.70,
    note="Сводные суммы раздела 1.", fill=TOTAL)
s1 = r
r += 1
for i, (m, a, b, va, vb) in enumerate([("Январь", 27093.00, 27093.00, 8127.90, 8127.90),
                                       ("Февраль", 27093.00, 27093.00, 8127.90, 8127.90),
                                       ("Март", 20101.26, 27093.00, 6030.38, 8127.90)]):
    row(ws, r, "  в т.ч. по месяцам (стр. 031–033)" if i == 0 else "", m, a, b, va, vb)
    r += 1
row(ws, r, "Контроль: раздел 1 − раздел 3", "", *[f"={L}{s1}-{L}{t3}" for L in "CDFG"],
    note="0 — внутри каждого расчёта раздел 1 сходится с разделом 3.", bold=True)
r += 2
text(ws, r, "Вывод по 1 кварталу:", [
    "1. Расхождение только по Лубковой Н.И. и только за март: выплаты +6 991,74 ₽, взносы +2 097,52 ₽ в уточнённом расчёте.",
    "2. Январь и февраль, а также Баскова Н.П. в обоих расчётах совпадают.",
    "3. В обоих расчётах взносы = 30% от выплат (арифметика верная).",
    "4. Доплата взносов за март 2 097,52 ₽ — срок уплаты был 28.04.2026, на эту сумму начисляются пени.",
    "5. Уточнённые расчёты (1 кв. и полугодие) поданы с КПП 366201001 в ИФНС 3662, первичные — с КПП 366301001 в ИФНС 3663.",
])

# ----------------------------------------------------------------- полугодие
ws, r = sheet("Полугодие", [
    "РСВ за полугодие 2026 г.: первичный и уточнённый — по сотрудникам",
    "ООО «Легкие деньки», ИНН 3663129244, тариф 01 (30%)",
    "Первичный: корр. 0 от 24.07.2026, КПП 366301001, ИФНС 3663.   Уточнённый: корр. 3 от 15.09.2026, КПП 366201001, ИФНС 3662.",
    "Разд. 3 РСВ за полугодие содержит только апрель–июнь. Январь–март по сотрудникам взяты из РСВ за 1 кв.: для первичного п/г — "
    "из первичного 1 кв., для уточнённого п/г — из уточнённого 1 кв. (суммы 1 кв. в разд. 1 совпадают, см. контроль ниже).",
], "первичный", "уточнённый")

r, h_l = person(ws, r, "Лубкова Наталья Ивановна", [
    ("Январь", 27093.00, 27093.00, 8127.90, 8127.90),
    ("Февраль", 27093.00, 27093.00, 8127.90, 8127.90),
    ("Март", 20101.26, 27093.00, 6030.38, 8127.90),
    ("Апрель", 0, 6321.70, 0, 1896.51),
], {"Март": "Как в РСВ за 1 кв.: первичный 20 101,26, уточнённый 27 093,00.",
    "Апрель": "В первичном Лубковой в апреле нет; в уточнённом добавлено 6 321,70 ₽."})
r, h_t = person(ws, r, "Троянова Екатерина Викторовна", [
    ("Апрель", 19868.20, 20771.30, 5960.46, 6231.39),
    ("Май", 27093.00, 27093.00, 8127.90, 8127.90),
    ("Июнь", 27093.00, 27093.00, 8127.90, 8127.90),
], {"Апрель": "Увеличено на 903,10 ₽. Апрель Лубкова + Троянова в уточнённом = 27 093,00 (МРОТ)."})
r, h_d = person(ws, r, "Денисова Александра Валерьевна", [
    ("Июнь", 21428.57, 21428.57, 6428.57, 6428.57),
], {"Июнь": "Принята в июне (30 000 × 15/21 раб. дн.)."})
r, h_b = person(ws, r, "Баскова Неонила Павловна", [
    ("Январь", 0, 0, 0, 0),
], {"Январь": "Выплат нет. В уточнённом п/г в стр. 010 не учтена (3 чел. вместо 4)."})

row(ws, r, "ИТОГО по сотрудникам", "п/г", *[f"={L}{h_l}+{L}{h_t}+{L}{h_d}+{L}{h_b}" for L in "CDFG"], fill=TOTAL, bold=True)
tot = r
r += 1
row(ws, r, "Раздел 1 РСВ за полугодие (стр. 030)", "п/г", 169770.03, 183986.57, 50931.01, 55195.97,
    note="Подр. 1 стр. 030 и разд. 1 стр. 030.", fill=TOTAL)
s1 = r
r += 1
row(ws, r, "Контроль: раздел 1 − сумма по сотрудникам", "", *[f"={L}{s1}-{L}{tot}" for L in "CDFG"], bold=True)
r += 2

ws.cell(r, 1, "Свод по кварталам и месяцам").font = Font(bold=True, size=11)
r += 1
for c, name in enumerate(COLS, 1):
    cell = ws.cell(r, c, name.format(a="первичный", b="уточнённый"))
    cell.font = Font(bold=True)
    cell.fill = HEAD
    cell.border = BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[r].height = 62
r += 1
row(ws, r, "1 квартал (п/г − стр. 031–033)", "янв–мар", f"=C{s1}-SUM(C{r+2}:C{r+4})", f"=D{s1}-SUM(D{r+2}:D{r+4})",
    f"=F{s1}-SUM(F{r+2}:F{r+4})", f"=G{s1}-SUM(G{r+2}:G{r+4})", note="Расхождение — Лубкова, март.")
q1r = r
r += 1
row(ws, r, "Контроль: 1 кв. в п/г − РСВ за 1 кв.", "", f"=C{q1r}-74287.26", f"=D{q1r}-81279", f"=F{q1r}-22286.18", f"=G{q1r}-24383.7",
    note="Первичный п/г сходится с первичным 1 кв., уточнённый п/г — с уточнённым 1 кв.", bold=True)
r += 1
for i, (m, a, b, va, vb, n) in enumerate([
        ("Апрель", 19868.20, 27093.00, 5960.46, 8127.90, "Расхождение: +6 321,70 Лубкова, +903,10 Троянова."),
        ("Май", 27093.00, 27093.00, 8127.90, 8127.90, "Совпадает."),
        ("Июнь", 48521.57, 48521.57, 14556.47, 14556.47, "Совпадает.")]):
    row(ws, r, "стр. 031–033" if i == 0 else "", m, a, b, va, vb, note=n)
    r += 1
row(ws, r, "Полугодие", "", f"=SUM(C{q1r},C{q1r+2}:C{r-1})", f"=SUM(D{q1r},D{q1r+2}:D{r-1})",
    f"=SUM(F{q1r},F{q1r+2}:F{r-1})", f"=SUM(G{q1r},G{q1r+2}:G{r-1})", fill=TOTAL, bold=True)
r += 2
text(ws, r, "Вывод по полугодию:", [
    "1. Уточнённый РСВ за полугодие больше первичного на 14 216,54 ₽ выплат и 4 264,96 ₽ взносов:",
    "   • Лубкова, март: +6 991,74 ₽ / +2 097,52 ₽ (то же, что в уточнённом РСВ за 1 кв.);",
    "   • Лубкова, апрель: +6 321,70 ₽ / +1 896,51 ₽ (в первичном не было);",
    "   • Троянова, апрель: +903,10 ₽ / +270,93 ₽ (20 771,30 вместо 19 868,20).",
    "2. Май, июнь, Денисова — без изменений. Взносы везде = 30% от выплат.",
    "3. Уточнённый РСВ за полугодие сходится с уточнённым РСВ за 1 кв. (1 кв. = 81 279,00 / 24 383,70).",
    "4. Доплата взносов: за март 2 097,52 ₽ (срок 28.04.2026), за апрель 2 167,44 ₽ (срок 28.05.2026) — на них начисляются пени.",
])

del wb["Sheet"]
for w in wb.worksheets:
    w.sheet_view.showGridLines = False
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth = 1
wb.save(OUT)
print("saved", OUT)
